from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .models import Position, Trade, position_from_dict, to_dict, trade_from_dict


@dataclass(slots=True)
class ImportResult:
    kind: str
    source: str
    row_count: int = 0
    imported: int = 0
    skipped: int = 0
    missing_columns: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    items: list[dict[str, Any]] = field(default_factory=list)


REQUIRED_COLUMNS = {
    "positions": {"symbol", "quantity", "cost_price", "current_price"},
    "trades": {"symbol", "action", "quantity", "price"},
}


def _normalize_header(header: str) -> str:
    return (
        str(header or "")
        .strip()
        .lower()
        .replace("\ufeff", "")
        .replace(" ", "_")
        .replace("-", "_")
        .replace("/", "_")
    )


FIELD_ALIASES = {
    "symbol": {
        "symbol",
        "code",
        "ticker",
        "ts_code",
        "stock_code",
        "security_code",
        "标的代码",
        "证券代码",
        "股票代码",
        "代码",
    },
    "name": {"name", "stock_name", "security_name", "标的名称", "证券名称", "股票名称", "名称"},
    "asset_type": {"asset_type", "asset type", "type", "资产类型", "类型"},
    "market": {"market", "exchange", "市场", "交易所", "板块"},
    "sector": {"sector", "industry", "行业", "行业板块"},
    "theme": {"theme", "topic", "concept", "主题", "概念"},
    "currency": {"currency", "ccy", "币种", "货币"},
    "quantity": {"quantity", "qty", "shares", "volume_holding", "数量", "持仓数量", "股数", "份额"},
    "cost_price": {"cost_price", "cost", "avg_cost", "average_cost", "成本价", "成本", "持仓成本", "平均成本"},
    "current_price": {"current_price", "last_price", "latest_price", "market_price", "当前价", "现价", "最新价", "市价"},
    "action": {"action", "side", "operation", "trade_type", "操作", "方向", "买卖方向", "交易类型"},
    "price": {"price", "trade_price", "成交价", "价格", "交易价格"},
    "amount": {"amount", "trade_amount", "成交金额", "金额", "交易金额"},
    "fee": {"fee", "commission", "手续费", "佣金", "费用"},
    "traded_at": {"traded_at", "trade_time", "trade_datetime", "成交时间", "交易时间", "时间"},
    "note": {"note", "remark", "memo", "备注", "说明"},
    "date": {"date", "trade_date", "交易日期", "日期"},
    "datetime": {"datetime", "trade_datetime", "trade_time", "交易时间", "时间", "日期时间"},
    "close": {"close", "close_price", "收盘", "收盘价"},
    "change_pct": {"change_pct", "pct_chg", "change_percent", "涨跌幅", "涨跌幅%"},
    "volume": {"volume", "vol", "成交量", "成交股数"},
    "period": {"period", "report_period", "end_date", "报告期", "期间", "财报期"},
    "metric": {"metric", "indicator", "指标", "科目"},
    "value": {"value", "metric_value", "数值", "值"},
}

ALIAS_LOOKUP = {
    _normalize_header(alias): canonical
    for canonical, aliases in FIELD_ALIASES.items()
    for alias in aliases
}

ACTION_ALIASES = {
    "buy": "buy",
    "b": "buy",
    "买": "buy",
    "买入": "buy",
    "加仓": "add",
    "add": "add",
    "reduce": "reduce",
    "减仓": "reduce",
    "sell": "sell",
    "s": "sell",
    "卖": "sell",
    "卖出": "sell",
    "清仓": "sell",
}


def import_records(path: str | Path, kind: str) -> ImportResult:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"Import file not found: {source}")
    headers, rows = read_tabular_rows(source)
    if kind == "positions":
        return _import_positions(source, headers, rows)
    if kind == "trades":
        return _import_trades(source, headers, rows)
    raise ValueError(f"Unsupported import kind: {kind}")


def read_tabular_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = [str(item or "").strip() for item in reader.fieldnames or []]
            return _normalize_table(headers, [dict(row) for row in reader])
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel_rows(path)
    raise ValueError(f"Unsupported import file type: {path.suffix}")


def _read_excel_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel import requires openpyxl.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(value or "").strip() for value in rows[0]]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        records.append(
            {
                headers[index]: row[index] if index < len(row) else None
                for index in range(len(headers))
                if headers[index]
            }
        )
    return _normalize_table(headers, records)


def _import_positions(
    source: Path,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> ImportResult:
    result = ImportResult(
        kind="positions",
        source=str(source),
        row_count=len(rows),
        missing_columns=_missing_columns(headers, "positions"),
    )
    if result.missing_columns:
        result.skipped = len(rows)
        result.errors.append(
            f"missing required columns: {', '.join(result.missing_columns)}"
        )
        return result
    for index, row in enumerate(rows, start=2):
        try:
            position = position_from_dict(
                {
                    "symbol": _text(row, "symbol").upper(),
                    "name": _text(row, "name") or _text(row, "symbol").upper(),
                    "asset_type": _text(row, "asset_type", "stock"),
                    "market": _text(row, "market", "A"),
                    "sector": _text(row, "sector", "unknown"),
                    "theme": _text(row, "theme", "unknown"),
                    "currency": _text(row, "currency", "CNY"),
                    "quantity": _number(row, "quantity"),
                    "cost_price": _number(row, "cost_price"),
                    "current_price": _number(row, "current_price"),
                }
            )
            if not position.symbol:
                raise ValueError("symbol is required")
            result.items.append(to_dict(position))
            result.imported += 1
        except Exception as exc:
            result.skipped += 1
            result.errors.append(f"row {index}: {exc}")
    return result


def _import_trades(
    source: Path,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> ImportResult:
    result = ImportResult(
        kind="trades",
        source=str(source),
        row_count=len(rows),
        missing_columns=_missing_columns(headers, "trades"),
    )
    if result.missing_columns:
        result.skipped = len(rows)
        result.errors.append(
            f"missing required columns: {', '.join(result.missing_columns)}"
        )
        return result
    for index, row in enumerate(rows, start=2):
        try:
            amount = _number(row, "amount")
            quantity = _number(row, "quantity")
            price = _number(row, "price")
            if amount == 0 and quantity and price:
                amount = quantity * price
            trade = trade_from_dict(
                {
                    "symbol": _text(row, "symbol").upper(),
                    "action": _normalize_action(_text(row, "action", "buy")),
                    "quantity": quantity,
                    "price": price,
                    "amount": amount,
                    "fee": _number(row, "fee"),
                    "traded_at": _text(row, "traded_at"),
                    "note": _text(row, "note"),
                }
            )
            if not trade.symbol:
                raise ValueError("symbol is required")
            result.items.append(to_dict(trade))
            result.imported += 1
        except Exception as exc:
            result.skipped += 1
            result.errors.append(f"row {index}: {exc}")
    return result


def _missing_columns(headers: list[str], kind: str) -> list[str]:
    normalized = {item.strip() for item in headers if item.strip()}
    return sorted(REQUIRED_COLUMNS[kind] - normalized)


def _normalize_table(
    headers: list[str],
    rows: list[dict[str, Any]],
) -> tuple[list[str], list[dict[str, Any]]]:
    normalized_headers = [_canonical_header(header) for header in headers]
    normalized_rows = [_normalize_row(row) for row in rows]
    return normalized_headers, normalized_rows


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for key, value in row.items():
        canonical = _canonical_header(str(key or ""))
        if not canonical:
            continue
        if canonical not in normalized or normalized[canonical] in (None, ""):
            normalized[canonical] = value
    return normalized


def _canonical_header(header: str) -> str:
    clean = _normalize_header(header)
    return ALIAS_LOOKUP.get(clean, clean)


def _text(row: dict[str, Any], key: str, default: str = "") -> str:
    value = row.get(key, default)
    if value is None:
        return default
    return str(value).strip()


def _number(row: dict[str, Any], key: str) -> float:
    value = row.get(key, 0)
    if value in (None, ""):
        return 0.0
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1]
    return float(text)


def _normalize_action(value: str) -> str:
    clean = str(value or "").strip().lower()
    return ACTION_ALIASES.get(clean, clean)
